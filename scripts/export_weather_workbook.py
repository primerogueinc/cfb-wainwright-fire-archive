"""Export all weather data in the archive as one Excel workbook plus one CSV per table, with every time shown
in UTC, local clock time (MDT, UTC-6) and, where a source uses it, Local Standard Time (LST, UTC-7).
Usage: python3 scripts/export_weather_workbook.py   (writes data/weather_exports/)"""
import os, re
from datetime import timedelta
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
E = os.path.join(REPO, "data/eccc"); OUT = os.path.join(REPO, "data/weather_exports"); os.makedirs(OUT, exist_ok=True)
H6, H7 = timedelta(hours=6), timedelta(hours=7)
HOT = PatternFill("solid", fgColor="FFE9B8")
sheets = []  # (sheet name, csv name, dataframe, time columns, highlight function, widths)


# 1-2. SWOB station observations
s = pd.read_csv(os.path.join(E, "swob_observations.csv"), dtype=str)
def swob(st):
    rows = []
    for _, r in s[s.tc_id == st].iterrows():
        z = pd.Timestamp(r.date_tm[:16]); g = r.max_wnd_spd_pst1hr_tm; gz = pd.NaT
        if isinstance(g, str) and g.strip():
            g = g.zfill(4); gz = z.replace(hour=int(g[:2]), minute=int(g[2:]))
            if gz > z: gz -= timedelta(days=1)
        num = lambda v, nd=1: None if pd.isna(v) else round(float(v), nd)
        rows.append({"obs_time_utc": z, "obs_time_mdt_local_clock_utc-6": z - H6, "obs_time_lst_utc-7": z - H7,
                     "station": r.tc_id, "station_name": r.stn_nam,
                     "air_temp_c": num(r.air_temp), "rel_hum_pct": num(r.rel_hum, 0), "dew_point_c": num(r.dwpt_temp),
                     "wind_dir_2min_avg_deg": num(r.avg_wnd_dir_10m_pst2mts, 0), "wind_spd_2min_avg_kmh": num(r.avg_wnd_spd_10m_pst2mts),
                     "wind_dir_10min_avg_deg": num(r.avg_wnd_dir_10m_pst10mts, 0), "wind_spd_10min_avg_kmh": num(r.avg_wnd_spd_10m_pst10mts),
                     "max_gust_past_hour_kmh": num(r.max_wnd_spd_10m_pst1hr), "max_gust_dir_deg": num(r.wnd_dir_10m_pst1hr_max_spd, 0),
                     "max_gust_time_utc": gz, "max_gust_time_mdt": gz - H6 if not pd.isna(gz) else pd.NaT,
                     "peak_wind_remark_raw": r.pk_wnd_rmk if isinstance(r.pk_wnd_rmk, str) else "",
                     "station_pressure_kpa": num(float(r.stn_pres) / 10, 2) if isinstance(r.stn_pres, str) else None,
                     "source_file": r.file.split("/")[-1]})
    return pd.DataFrame(rows)
T_SW = ["obs_time_utc", "obs_time_mdt_local_clock_utc-6", "obs_time_lst_utc-7", "max_gust_time_utc", "max_gust_time_mdt"]
cz = swob("CZWN"); ll = swob("CYLL")
sheets.append(("CZWN obs Oct 17", "1_czwn_swob_oct17.csv", cz, T_SW, lambda r: str(r["peak_wind_remark_raw"]).endswith("2116Z")))
sheets.append(("CYLL obs Oct 17", "2_cyll_swob_oct17.csv", ll, T_SW, lambda r: r["obs_time_utc"] == pd.Timestamp("2026-10-17 22:00")))

# 3. Climate hourly (archive, LST)
c = pd.read_csv(os.path.join(E, "climate_hourly/wain_hourly_2026-10.csv"))
lst = pd.to_datetime(c["Date/Time (LST)"])
c.insert(0, "time_lst_as_published_utc-7", lst); c.insert(0, "time_mdt_local_clock_utc-6", lst + timedelta(hours=1)); c.insert(0, "time_utc", lst + H7)
c = c.drop(columns=["Date/Time (LST)", "Year", "Month", "Day", "Time (LST)"])
c["wind_dir_deg"] = c["Wind Dir (10s deg)"] * 10
c = c.dropna(axis=1, how="all")
sheets.append(("Climate hourly Oct 2026", "3_climate_hourly_czwn_oct2026.csv", c, ["time_utc", "time_mdt_local_clock_utc-6", "time_lst_as_published_utc-7"],
               lambda r: pd.Timestamp("2026-10-17 13:00") <= r["time_mdt_local_clock_utc-6"] <= pd.Timestamp("2026-10-17 17:00")))

# 4. Forecasts and warnings naming Wainwright, with the Wainwright text
b = pd.read_csv(os.path.join(E, "bulletins_index.csv"))
rows = []
for _, r in b[(b.mentions_wainwright) & (~b.duplicate_copy)].iterrows():
    t = open(os.path.join(E, r.file), errors="replace").read()
    m = re.search(r"\bAT\s+(\d{1,2}:\d{2}\s+[AP]\.M\.\s+M[DS]T)", t) or re.search(r"\bA\s+(\d{1,2}H\d{2}\s+HA[RN])", t)
    if r.kind == "WW":
        body = t.split("\n", 1)[1].strip()
    else:
        blk = [x for x in re.split(r"\n\s*\n", t) if "WAINWRIGHT" in x.upper()]
        body = blk[0].strip() if blk else ""
    rows.append({"issued_utc": pd.Timestamp(r.issued_utc), "issued_mdt": pd.Timestamp(r.issued_mdt),
                 "time_printed_on_bulletin": m.group(1) if m else "", "bulletin": r.header, "office": r.office,
                 "type": {"FP": "Forecast", "WW": "Warning"}.get(r.kind, r.kind),
                 "language": "French" if r.header.endswith(("76", "58")) else "English",
                 "wainwright_text": body, "source_file": r.file})
fw = pd.DataFrame(rows)
sheets.append(("Wainwright forecasts+warnings", "4_wainwright_forecasts_warnings.csv", fw, ["issued_utc", "issued_mdt"],
               lambda r: r["bulletin"] in ("WWCN16",) or (r["bulletin"] == "FPCN16" and r["issued_utc"] == pd.Timestamp("2026-10-17 10:30"))))

# 5. All bulletins
ball = b.copy(); ball["issued_utc"] = pd.to_datetime(ball.issued_utc); ball["issued_mdt"] = pd.to_datetime(ball.issued_mdt)
sheets.append(("All bulletins Oct 16-18", "5_all_bulletins_index.csv", ball, ["issued_utc", "issued_mdt"], lambda r: bool(r["mentions_wainwright"])))

# 6. Fire Weather Index
f = pd.read_csv(os.path.join(REPO, "data/nrcan_cwfis/fwi_nearby_stations_oct10-18.csv"), skipinitialspace=True)
f.columns = [x.strip() for x in f.columns]
for col in f.columns:
    if f[col].dtype == object: f[col] = f[col].str.strip()
d = pd.to_datetime(f.REPDATE.astype(str), format="%Y%m%d")
f.insert(0, "weather_time_mdt", d + timedelta(hours=13)); f.insert(0, "weather_time_lst_noon_utc-7", d + timedelta(hours=12)); f.insert(0, "date", d)
f = f.replace("NULL", None)
sheets.append(("Fire Weather Index", "6_fwi_nearby_stations_oct10-18.csv", f, ["date", "weather_time_lst_noon_utc-7", "weather_time_mdt"],
               lambda r: r["date"] == pd.Timestamp("2026-10-17") and "WAINWRIGHT" in str(r["NAME"]).upper()))

README = [
    "CFB Wainwright fire, Oct. 17, 2026: weather data in the Prime Rogue archive, as one workbook.",
    "Compiled by Prime Rogue Inc. from public data. Each sheet is also posted as a CSV with the same content.",
    "",
    "TIME ZONES. On Oct. 17, 2026 local clock time in Wainwright was MDT = UTC - 6. 21:16 UTC = 3:16 p.m.",
    "Some sources use Local Standard Time, LST = UTC - 7, all year. LST hours read one hour EARLIER than the clock that day.",
    "Every time column is labelled with its zone. Times in a source's own zone are kept, and the others are calculated from it.",
    "",
    "SHEETS",
    "CZWN obs Oct 17: base weather station (Wainwright CFB Airfield 21), hourly SWOB observations from the MSC Datamart. Source times are UTC. "
    "Each row's gust columns give the strongest gust in the hour BEFORE that observation: the 82 km/h gust at 21:16 UTC (3:16 p.m. MDT) is in the 22:00 UTC row (highlighted). "
    "Rows at 02, 10, 12 and 16 UTC were not in the pull and are not filled in.",
    "CYLL obs Oct 17: Lloydminster airport, the three SWOB hours pulled (2-4 p.m. MDT). Gust time not reported in these files.",
    "Climate hourly Oct 2026: Environment Canada historical climate archive, station 301S001, all of October. Published in LST (UTC - 7). "
    "time_utc and time_mdt are added. Hourly means only, no gusts. Oct. 17, 1-5 p.m. MDT highlighted. Matches the SWOB hour for hour once the offset is applied.",
    "Wainwright forecasts+warnings: every public forecast (FPCN16 English, FPCN76 French, FPCN54/58 extended) and wind warning (WWCN16/76) naming Wainwright, Oct. 16-18. "
    "issued_utc comes from the bulletin's file name. time_printed_on_bulletin is the time Environment Canada prints in the text. Scheduled forecasts carry the scheduled time, "
    "which can differ from transmission: the forecast transmitted at 4:30 a.m. MDT on Oct. 17 (10:30 UTC) is printed as '5:00 A.M. MDT' (highlighted, with the warnings).",
    "All bulletins Oct 16-18: index of every bulletin retrieved from the Prairie and Northern storm-prediction offices.",
    "Fire Weather Index: NRCan CWFIS daily values for the stations nearest the base. Computed from noon LST observations = 1 p.m. MDT. Alberta classes: Extreme at FWI 29.5 and above.",
    "",
    "SOURCES. Environment and Climate Change Canada (MSC Datamart; Historical Climate Data). Natural Resources Canada (CWFIS). Open Government Licence - Canada.",
    "Originals and scripts: https://github.com/primerogueinc/cfb-wainwright-fire-archive (branch authentic-stoic-alternative-hypothesis), data/eccc and data/nrcan_cwfis.",
]

wb = Workbook(); ws = wb.active; ws.title = "Read me"
for line in README: ws.append([line])
ws.column_dimensions["A"].width = 160; ws["A1"].font = Font(bold=True)
for row in ws.iter_rows():
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")

for name, csvn, df, tcols, hot, *_ in sheets:
    out = df.copy()
    for t in tcols: out[t] = pd.to_datetime(out[t])
    csv = out.copy()
    for t in tcols: csv[t] = csv[t].dt.strftime("%Y-%m-%d" if t == "date" else "%Y-%m-%d %H:%M").fillna("")
    csv.to_csv(os.path.join(OUT, csvn), index=False)
    ws = wb.create_sheet(name[:31]); cols = list(out.columns); ws.append(cols)
    for _, r in out.iterrows():
        vals = []
        for cname in cols:
            v = r[cname]
            if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
            elif v is pd.NaT or (not isinstance(v, str) and pd.isna(v)): v = None
            elif hasattr(v, "item"): v = v.item()
            vals.append(v)
        ws.append(vals)
        if hot(r):
            for cell in ws[ws.max_row]: cell.fill = HOT
    for i, cname in enumerate(cols, 1):
        L = get_column_letter(i)
        if cname in tcols:
            w = 18
            for cell in ws[L][1:]: cell.number_format = "yyyy-mm-dd" if cname == "date" else "yyyy-mm-dd hh:mm"
        elif cname == "wainwright_text": w = 90
        elif cname in ("source_file", "file", "headlines", "station_name", "Station Name", "NAME"): w = 30
        else: w = max(10, min(22, len(str(cname)) * 0.9))
        ws.column_dimensions[L].width = w
        if cname == "wainwright_text":
            for cell in ws[L][1:]: cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]: cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45; ws.freeze_panes = "B2"
wb.save(os.path.join(OUT, "cfb-wainwright-weather-data-oct2026.xlsx"))
open(os.path.join(OUT, "README.txt"), "w").write("\n".join(README) + "\n\nCSV files: " + ", ".join(s[1] for s in sheets) + "\n")
for name, csvn, df, *_ in sheets: print(f"{name}: {len(df)} rows -> {csvn}")
